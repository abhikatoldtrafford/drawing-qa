"""Benchmark a generated BOQ against a reference BOQ workbook (e.g. the user's 2GU1 BOQ.xlsx).

Compares, row by row (matched on ITEMNO): every value column, the row-1 SUBTOTALs, and the formula
structure of the exported sheet (J = H*I, L = W*L*QTY*UNIT or L*QTY*UNIT, M = L*I, O = N*I, P = M-O)."""
import io
import re

import openpyxl

from .export import boq_excel

COLUMNS = [("A", "drawing_no"), ("B", "item_type"), ("C", "mark_no"), ("D", "item_no"), ("E", "section"),
           ("F", "width"), ("G", "length"), ("H", "qty"), ("I", "fab_qty"), ("J", "total_qty"), ("K", "unit_wt"),
           ("L", "calc_wt"), ("M", "total_calc_wt"), ("N", "drg_wt"), ("O", "total_drg_wt"), ("P", "difference"),
           ("Q", "grade")]
FORMULA_COLS = "JLMOP"


def _rows(ws, first=3):
    out = {}
    for r in range(first, ws.max_row + 1):
        item = ws[f"D{r}"].value
        if item not in (None, ""):
            out[str(item)] = r
    return out


def _same(a, b, tol):
    if isinstance(a, (int, float)) or isinstance(b, (int, float)):
        if a is None or b is None:
            return a is None and b is None
        return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))
    return (a or "") == (b or "")


def _shape(formula, row):
    """'=(F3/1000)*(G3/1000)*H3*K3' -> '=(F#/1000)*(G#/1000)*H#*K#' so formulas compare across rows."""
    return re.sub(rf"([A-Z]){row}\b", r"\1#", str(formula)) if isinstance(formula, str) else formula


def compare_boq(inventory, gt_path, tol=1e-6) -> dict:
    gt_v = openpyxl.load_workbook(gt_path, data_only=True).active
    gt_f = openpyxl.load_workbook(gt_path, data_only=False).active
    ours_f = openpyxl.load_workbook(io.BytesIO(boq_excel(inventory)), data_only=False).active
    ours = {r.item_no: r for r in inventory.boq}
    gt_rows, our_rows = _rows(gt_v), _rows(ours_f)
    diffs, formula_diffs, checked = [], [], 0
    for item, gr in gt_rows.items():
        r = ours.get(item)
        if r is None:
            continue
        for col, attr in COLUMNS:
            checked += 1
            if not _same(getattr(r, attr), gt_v[f"{col}{gr}"].value, tol):
                diffs.append({"item": item, "field": attr, "ours": getattr(r, attr), "gt": gt_v[f"{col}{gr}"].value})
        orow = our_rows[item]
        for col in FORMULA_COLS:
            if _shape(gt_f[f"{col}{gr}"].value, gr) != _shape(ours_f[f"{col}{orow}"].value, orow):
                formula_diffs.append({"item": item, "col": col, "ours": ours_f[f"{col}{orow}"].value,
                                      "gt": gt_f[f"{col}{gr}"].value})
    totals = {}
    for col, attr in (("J", "total_qty"), ("L", "calc_wt"), ("M", "total_calc_wt"), ("N", "drg_wt"),
                      ("O", "total_drg_wt"), ("P", "difference")):
        mine = sum(getattr(r, attr) or 0 for r in inventory.boq)
        totals[attr] = {"ours": mine, "gt": gt_v[f"{col}1"].value, "match": _same(mine, gt_v[f"{col}1"].value, 1e-6)}
    return {
        "gt_rows": len(gt_rows), "our_rows": len(ours),
        "missing": sorted(set(gt_rows) - set(ours)), "extra": sorted(set(ours) - set(gt_rows)),
        "fields_checked": checked, "field_diffs": diffs, "formula_diffs": formula_diffs, "totals": totals,
    }
