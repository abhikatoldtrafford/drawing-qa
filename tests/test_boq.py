import io

import openpyxl
import pytest

from drawing_qa import export
from drawing_qa.boq import build_boq, item_type, missing_unit_weights
from drawing_qa.inventory import deterministic_inventory
from drawing_qa.steel_tables import handbook_unit_weight, plate_unit_weight

from .conftest import GOLDEN, SAMPLES, det, sample_path

# The user's reference BOQ for 16807 (2GU1 BOQ.xlsx): item, section, width, length, qty, unit wt, calculated wt
REFERENCE_16807 = [
    ("2GU1", "PL6", 1067, 1478.4, 1, 47.1, 74.29802688),
    ("2m272", "ISMC150", None, 9088, 1, 16.8, 152.6784),
    ("2m273", "ISA75X75X8", None, 9088, 1, 8.9, 80.8832),
    ("2p205", "PL8", 50, 400.3, 18, 62.8, 22.624956),
    ("2p206", "PL8", 50, 738.6, 18, 62.8, 41.745672),
    ("2p207", "PL8", 50, 451.3, 18, 62.8, 25.507476),
    ("2p208", "PL6", 1478, 2000, 4, 47.1, 556.9104),
    ("2p209", "PL6", 75, 1412.4, 5, 47.1, 24.946515),
]


def test_handbook_unit_weights():
    assert handbook_unit_weight("ISMC150") == (16.8, "IS 808 table")
    assert handbook_unit_weight("ISA75X75X8") == (8.9, "IS 808 table")
    assert handbook_unit_weight("ISA 200X200X20")[0] == 60.0
    assert handbook_unit_weight("ISA50X50X6")[0] == 4.5 and handbook_unit_weight("ISMC300")[0] == 36.3
    assert handbook_unit_weight("ROD20")[0] == 2.47
    assert handbook_unit_weight("ROD8")[0] == pytest.approx(0.395, abs=1e-3)          # computed, not tabulated
    assert handbook_unit_weight("PIPE508*6")[0] == pytest.approx(74.28, abs=0.01)
    assert handbook_unit_weight("SPD508*508*608*608*8") == (None, "")
    assert plate_unit_weight(6) == pytest.approx(47.1)


@pytest.mark.parametrize("drg,expected", [("09970", "DOWN COMER"), ("14281", "COLUMN"), ("16807", "GUTTER")])
def test_item_type_from_title(drg, expected, settings):
    assert item_type(det(drg, settings)) == expected


def test_boq_16807_matches_reference_rows(settings):
    rows = {r.item_no: r for r in build_boq(det("16807", settings))}
    assert list(rows) == [r[0] for r in REFERENCE_16807]
    for item, section, width, length, qty, unit, calc in REFERENCE_16807:
        r = rows[item]
        assert (r.section, r.width, r.length, r.qty, r.fab_qty, r.total_qty) == (section, width, length, qty, 1, qty)
        assert r.unit_wt == pytest.approx(unit) and r.calc_wt == pytest.approx(calc, abs=1e-6), item
        assert (r.drawing_no, r.item_type, r.mark_no) == ("TST-SFD-46-01-01-07-000-16807", "GUTTER", "2GU1")
    assert sum(r.difference for r in rows.values()) == pytest.approx(2.75464588, abs=1e-6)


def test_fab_qty_multiplies_totals_16362(settings):
    rows = {r.item_no: r for r in build_boq(det("16362", settings))}
    rod = rows["1m473"]
    assert (rod.qty, rod.fab_qty, rod.total_qty) == (22, 2, 44)
    assert rod.total_calc_wt == pytest.approx(2 * rod.calc_wt) and rod.total_drg_wt == pytest.approx(2 * 4.75)


def test_built_ups_are_listed_as_their_plates_14281(settings):
    rows = {r.item_no: r for r in build_boq(det("14281", settings))}
    fl, web = rows["3m471 flange"], rows["3m471 web"]
    assert (fl.section, fl.width, fl.qty, web.section, web.width, web.qty) == ("PL40", 500, 2, "PL32", 1120, 1)
    assert fl.drg_wt + web.drg_wt == pytest.approx(4518.36, abs=0.01)             # BOM gross of 3m471 split
    assert abs(fl.difference) < 0.01 and abs(web.difference) < 0.01


@pytest.mark.parametrize("drg", sorted(GOLDEN))
def test_boq_drawing_weight_reconciles_on_every_sheet(drg, settings):
    inv = deterministic_inventory(det(drg, settings))
    assert {c.id: c.level for c in inv.checks}["boq.drawing_weight"] == "pass"
    assert sum(r.total_drg_wt for r in inv.boq) == pytest.approx(GOLDEN[drg]["gross"], abs=0.1)


def test_section_without_handbook_value_uses_drawing_weight_and_is_flagged(settings):
    x = det("16362", settings)
    assert missing_unit_weights(x) == {"SPD508*508*608*608*8": pytest.approx(54.93 / 0.506, abs=1e-3)}
    inv = deterministic_inventory(x)
    spd = next(r for r in inv.boq if r.item_no == "1m470")
    assert spd.unit_wt_source.startswith("drawing") and "verify" in spd.note
    assert "1m470" in next(c for c in inv.checks if c.id == "boq.calculated").message


def test_boq_excel_has_reference_layout_and_live_formulas(settings):
    inv = deterministic_inventory(det("16807", settings))
    ws = openpyxl.load_workbook(io.BytesIO(export.boq_excel(inv))).active
    assert [c.value for c in ws[2]][:17] == export.BOQ_HEADERS[:17]
    assert ws["J1"].value == "=SUBTOTAL(9,J3:J10)" and ws["P1"].value == "=SUBTOTAL(9,P3:P10)"
    assert (ws["J3"].value, ws["K3"].value, ws["L3"].value) == ("=H3*I3", "=7.85*6", "=(F3/1000)*(G3/1000)*H3*K3")
    assert (ws["K4"].value, ws["L4"].value, ws["P4"].value) == (16.8, "=(G4/1000)*H4*K4", "=M4-O4")


def test_benchmark_against_reference_workbook(settings):
    gt = SAMPLES / "2GU1 BOQ.xlsx"
    if not gt.exists():
        pytest.skip("reference BOQ 2GU1 BOQ.xlsx not present")
    from drawing_qa.benchmark import compare_boq
    res = compare_boq(deterministic_inventory(det("16807", settings)), gt)
    assert (res["our_rows"], res["gt_rows"], res["missing"], res["extra"]) == (8, 8, [], [])
    # the only difference is the reference's own typo: grade 'E2350A' for 2GU1 (the drawing's BOM says E350A)
    assert res["field_diffs"] == [{"item": "2GU1", "field": "grade", "ours": "E350A", "gt": "E2350A"}]
    assert res["formula_diffs"] == []
    assert all(v["match"] for v in res["totals"].values())
