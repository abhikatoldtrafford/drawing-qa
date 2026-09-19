import io
import json

import openpyxl

from drawing_qa import export

from .conftest import det


def test_exports(settings):
    from drawing_qa.inventory import deterministic_inventory
    x = det("14281", settings)
    base = ["Title", "BOM", "Locations", "Abstract", "Bolts", "Revisions", "Notes", "Views", "Checks"]
    assert openpyxl.load_workbook(io.BytesIO(export.to_excel(x))).sheetnames == base
    wb = openpyxl.load_workbook(io.BytesIO(export.to_excel(x, deterministic_inventory(x))))
    assert wb.sheetnames == ["BOQ"] + base + ["Inv Summary", "Inv Plates", "Inv Sections", "Inv Fasteners", "Inv Welds",
                                    "Inv Review", "Inv Checks"]
    assert wb["Inv Plates"].max_row == 9 + 1
    assert wb["BOM"].max_row == 71 + 1
    csv = export.bom_csv(x).splitlines()
    assert csv[0].startswith("item_no,section,length_mm,qty") and len(csv) == 72
    assert json.loads(export.to_json(x))["title_block"]["drawing_no"].endswith("14281")
